"""Create fresh test DB, seed users from users.xlsx, import vendors."""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

os.environ["SC_GR_DEV"] = "1"

from sc_gr_app.config import default_config
from sc_gr_app.db.connection import connect
from sc_gr_app.db.migrations import migrate

BASE = os.path.dirname(os.path.abspath(__file__))
VENDOR_XLSX = os.path.join(BASE, "Vendors_2026-07-07.xlsx")
USERS_XLSX = os.path.join(BASE, "users.xlsx")

# -- 1. Delete old DB --
config = default_config()
if config.db_path.exists():
    config.db_path.unlink()
    print(f"Deleted old DB: {config.db_path}")

# -- 2. Create fresh schema --
migrate(config)
print("Created fresh schema.")

# -- 3. Seed users from users.xlsx --
import openpyxl
wb = openpyxl.load_workbook(USERS_XLSX)
ws = wb["Tabelle1"]

from datetime import datetime, timezone
ts = datetime.now(timezone.utc).isoformat()

users_added = 0
with connect(config) as conn:
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        machine_id, name, email, role = row[0], row[1], row[2], row[3]
        if not machine_id:
            continue
        machine_id = str(machine_id).strip()
        name = str(name).strip() if name else ""
        email = str(email).strip() if email else ""
        role = str(role).strip() if role else "requester"

        user_id = f"U-{machine_id}"
        conn.execute(
            """INSERT OR IGNORE INTO users (
              user_id, machine_id, user_name, role, email, status, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, 'active', ?, ?)""",
            (user_id, machine_id, name, role, email, ts, ts),
        )
        users_added += 1
    conn.commit()
print(f"Seeded {users_added} users from users.xlsx.")

# -- 4. Service scope normalization --
SCOPE_FIXES = {
    "Secirity": "Security",
    "engineering Service": "Engineering Service",
    "Materials": "Others",
    "Driver/Test car rental": "Driver",
    "insurance": "Insurance",
}

SUPPORTED_SCOPES = {
    "Transportation", "Engineering Service", "Equipment", "Parts",
    "Driver", "Test car rental", "General Service", "Dealers",
    "Import&Export&cusoms clearance", "Insurance", "Harness",
    "Maintenance&Calibration", "Security", "Testing support",
    "Fix asset", "Materials", "Others",
}

# Get admin user for created_by
with connect(config) as conn:
    admin = dict(conn.execute(
        "SELECT * FROM users WHERE role = 'admin' LIMIT 1"
    ).fetchone())
    # Verify users
    users = conn.execute("SELECT user_id, user_name, role FROM users ORDER BY role, user_name").fetchall()
    print(f"\n=== Users in DB ({len(users)}) ===")
    for u in users:
        print(f"  {u['user_id']}: {u['user_name']} [{u['role']}]")

# -- 5. Import vendors from xlsx --
# Clear existing data from seed DB (keep users and schema)
with connect(config) as conn:
    conn.execute("PRAGMA foreign_keys = OFF")
    gr_deleted = conn.execute("DELETE FROM gr_requests").rowcount
    po_deleted = conn.execute("DELETE FROM pos").rowcount
    sc_deleted = conn.execute("DELETE FROM sc_records").rowcount
    v_deleted = conn.execute("DELETE FROM vendors").rowcount
    conn.commit()
    print(f"Cleared seed DB data: {sc_deleted} SC, {po_deleted} PO, {gr_deleted} GR, {v_deleted} vendors")

imported = 0
skipped = 0
errors = []

wb = openpyxl.load_workbook(VENDOR_XLSX)
ws = wb["Sheet1"]
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
    _vid = str(row[0]).strip() if row[0] else ""
    vendor_name = str(row[1] or "").strip()
    company_cn = str(row[2] or "").strip()
    ksrm_code = str(row[3] or "").strip() if row[3] else ""
    scope = str(row[4] or "").strip()
    contact = str(row[5] or "").strip() if row[5] else ""
    phone = str(row[6] or "").strip() if row[6] else ""
    email = str(row[7] or "").strip() if row[7] else ""
    vid = _vid

    if not vendor_name or not scope:
        skipped += 1
        errors.append(f"Row {i+1}: missing name or scope, skipped")
        continue

    if scope in SCOPE_FIXES:
        print(f"  [!] Fixing scope: '{scope}' -> '{SCOPE_FIXES[scope]}' ({vendor_name})")
        scope = SCOPE_FIXES[scope]

    if scope not in SUPPORTED_SCOPES:
        skipped += 1
        errors.append(f"Row {i+1}: unsupported scope '{scope}' for '{vendor_name}', skipped")
        continue

    with connect(config) as conn:
        if not vid or vid == vendor_name:
            row_max = conn.execute(
                "SELECT COALESCE(MAX(CAST(SUBSTR(vendor_id, 2) AS INTEGER)), 0) + 1 AS n FROM vendors"
            ).fetchone()
            vid = f"V{row_max['n']:06d}"

        try:
            conn.execute(
                """INSERT INTO vendors (
                  vendor_id, vendor_name, ksrm_vendor_code, company_name_cn,
                  contact_person, phone, service_scope, email,
                  description, inquiry_history, created_by, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    vid,
                    vendor_name,
                    ksrm_code or None,
                    company_cn or None,
                    contact or None,
                    phone or None,
                    scope,
                    email or None,
                    None,
                    None,
                    admin["user_id"],
                    ts, ts,
                ),
            )
            conn.commit()
            imported += 1
        except Exception as e:
            skipped += 1
            errors.append(f"Row {i+1} '{vendor_name}': {e}")
            continue

print(f"\n=== Vendor Import ===")
print(f"  Imported: {imported}")
print(f"  Skipped:  {skipped}")
if errors:
    print(f"  Errors:")
    for e in errors:
        print(f"    - {e}")

# Verify
with connect(config) as conn:
    v_count = conn.execute("SELECT COUNT(*) as n FROM vendors").fetchone()["n"]
    u_count = conn.execute("SELECT COUNT(*) as n FROM users").fetchone()["n"]
print(f"\nDB state: {u_count} users, {v_count} vendors")
print("Done.")
