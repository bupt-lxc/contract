"""
Extract data from Internal request collection xlsx -> full_data.csv + import files.
Updated with vendor fixes, placeholder PO skip, Internal System Number handling,
PO amount lookup from related rows, and contract type rules.
"""
import csv
import os
import sys
from collections import defaultdict

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "Internal request collection_CW27.4&28.4_2026_SJL - final.xlsx")
VENDOR_XLSX = os.path.join(BASE, "Vendors_2026-07-07.xlsx")
USERS_XLSX = os.path.join(BASE, "users.xlsx")

# ── Vendor name fixes (data -> correct DB name) ───────────────────────
VENDOR_NAME_FIXES = {
    "Beijing CEFOC-CARE facility management co. Ltd. 3300239993":
        "BeijingCEFOC-CAREfacilitymanagementco.Ltd",
    "CTS International Logistics Corporation":
        "CTS International Logistics Corporation Limited-Beijingbranch",
    "China Security & Protection Group":
        "China Security & Protection Group Co.,Ltd",
    "Shanghai Zhongcheng Automotive Engineering Co., Limited":
        "ShanghaiZhongchengAutomotive",
}

# ── Load xlsx ──────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["C26.4"]

# Headers from row 2
xlsx_headers = []
for col in range(1, 37):
    val = ws.cell(row=2, column=col).value
    xlsx_headers.append(str(val).replace("\n", " ").strip() if val else f"COL_{col}")

# Column index helper
def col_idx(name_substring):
    for i, h in enumerate(xlsx_headers):
        if name_substring.lower() in h.lower():
            return i
    return -1

# ── Extract all data rows ──────────────────────────────────────────────
data_rows = []
for r in range(3, 288):
    row = {}
    for col in range(1, 37):
        val = ws.cell(row=r, column=col).value
        row[xlsx_headers[col - 1]] = str(val).strip() if val is not None else ""
    if any(v for v in row.values()):
        row["_excel_row"] = r
        data_rows.append(row)

print(f"Extracted {len(data_rows)} data rows")

# ── Build requester name -> user_id mapping ────────────────────────────
wb_users = openpyxl.load_workbook(USERS_XLSX, data_only=True)
ws_users = wb_users["Tabelle1"]
requester_map = {}  # normalized_name -> user_id
for row in ws_users.iter_rows(min_row=2, values_only=True):
    machine_id, name, email, role = row[0], row[1], row[2], row[3]
    if not machine_id:
        continue
    mid = str(machine_id).strip()
    user_id = f"U-{mid}"
    full_name = str(name).strip() if name else ""
    if full_name:
        # Normalize: remove parenthetical suffix like "(C/EV-L)", remove commas
        short = full_name.split("(")[0].strip().rstrip()
        requester_map[short] = user_id
        requester_map[full_name] = user_id
        # Also match without comma: "Zhou Liwei" matches "Zhou, Liwei"
        no_comma = short.replace(",", "")
        requester_map[no_comma] = user_id
        requester_map[short.replace(",", " ").strip()] = user_id
print(f"Built {len(set(requester_map.values()))} requester name mappings")

# Apply vendor name fixes + filter empty vendors
fixed_vendor = 0
skipped_empty_vendor = 0
for row in data_rows:
    v = row["Vendor"]
    if not v:
        skipped_empty_vendor += 1
        row["_skip"] = True
        continue
    if v in VENDOR_NAME_FIXES:
        row["Vendor"] = VENDOR_NAME_FIXES[v]
        fixed_vendor += 1
        print(f"  Fixed vendor: '{v}' -> '{row['Vendor']}'")

if fixed_vendor:
    print(f"Fixed {fixed_vendor} vendor names")
if skipped_empty_vendor:
    print(f"Skipped {skipped_empty_vendor} rows with empty vendor")

# ── Build lookup maps for related-row data ─────────────────────────────
# For each PO No, find the SC Approved or PO Approved row that has PO amount
po_amount_map = {}  # po_no -> po_amount from SC Approved/PO Approved row
sc_amount_map = {}  # sc_no -> sc_amount
po_contract_from_map = {}
po_contract_to_map = {}
for row in data_rows:
    s = row["Status"]
    po_no = row["PO No."]
    sc_no = row["SC No."]
    if s in ("SC Approved", "PO Approved"):
        if po_no:
            po_amt = row["PO amount                  (RMB)"]
            if po_amt and po_no not in po_amount_map:
                po_amount_map[po_no] = po_amt
            cf = row["Contract from  mm/dd/yyyy"] or row["CON Date"]
            ct = row["Contract to  mm/dd/yyyy"]
            if cf and po_no not in po_contract_from_map:
                po_contract_from_map[po_no] = cf
            if ct and po_no not in po_contract_to_map:
                po_contract_to_map[po_no] = ct
    if s in ("SC Approved", "SC Pending") and sc_no:
        sc_amt = row["SC amount (RMB) incl. VAT"]
        if sc_amt and sc_no not in sc_amount_map:
            sc_amount_map[sc_no] = sc_amt

# ── Step 1: Write full_data.csv ────────────────────────────────────────
FD_FIELDS = [
    "service_scope", "year", "pos", "sc_no", "po_no",
    "contract_type", "status", "vendor", "description",
    "net_gr", "gr_vat", "con_val", "open_po", "po_amt", "sc_amt",
    "category", "req_raw", "cost_center", "req_date", "conf_no",
    "con_date", "conf_from", "conf_to", "pay_freq", "purchaser",
    "supplier", "asset", "remark", "asset_num",
    "record_type", "action", "problem", "auto_fix",
    "sc_import_error", "po_import_error", "gr_import_error",
]

full_rows = []
for row in data_rows:
    fd = {k: "" for k in FD_FIELDS}
    fd["service_scope"] = row["Service scope"]
    fd["year"] = row["BG Year"]
    fd["pos"] = row["Contract  Pos."]
    fd["sc_no"] = row["SC No."]
    fd["po_no"] = row["PO No."]
    fd["contract_type"] = row["Type of Contract"]
    fd["status"] = row["Status"]
    fd["vendor"] = row["Vendor"]
    fd["description"] = row["Description"]
    fd["net_gr"] = row["NET GR Application Amount"]
    fd["gr_vat"] = row["GR  Application amount incl. VAT"]
    fd["con_val"] = row["CON Value"]
    fd["open_po"] = row["Open PO Value incl. VAT"]
    fd["po_amt"] = row["PO amount                  (RMB)"]
    fd["sc_amt"] = row["SC amount (RMB) incl. VAT"]
    fd["category"] = row["Category"]
    fd["req_raw"] = row["Requestor"]
    fd["cost_center"] = row["Cost Center"]
    fd["req_date"] = row["Request date"]
    fd["conf_no"] = row["Confirmation No."]
    fd["con_date"] = row["CON Date"]
    fd["conf_from"] = row["Contract from  mm/dd/yyyy"]
    fd["conf_to"] = row["Contract to  mm/dd/yyyy"]
    fd["pay_freq"] = row["Payment Frequency"]
    fd["purchaser"] = row["Purchaser"]
    fd["supplier"] = row["Supplier code (Option)"]
    fd["asset"] = row["Asset (Y/N)"]
    fd["remark"] = row["Remark"]
    fd["asset_num"] = row["Asset number"]
    fd["record_type"] = row["Status"]
    if row.get("_skip"):
        fd["action"] = "Skipped"
        fd["problem"] = "Vendor is empty"
    else:
        fd["action"] = "Imported"
    full_rows.append(fd)

with open(os.path.join(BASE, "full_data.csv"), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=FD_FIELDS)
    w.writeheader()
    w.writerows(full_rows)
print(f"Wrote {len(full_rows)} rows to full_data.csv")

# ── Determine request type and call_off handling ───────────────────────
# Internal System Number -> call_off rule (only for PO Approved / SC Approved)
def is_call_off(row):
    """Check if this is a call-off SC/PO based on Internal System Number."""
    isn = row["Internal System Number （有号写号，没号写N/A）"]
    return isn and isn.upper() != "N/A"


# ── Build explicit SC/PO indexes ───────────────────────────────────────
explicit_sc = set()
explicit_po = set()
for row in data_rows:
    if row.get("_skip"):
        continue
    s = row["Status"]
    sc_no = row["SC No."]
    po_no = row["PO No."]
    # Don't include placeholder PO numbers
    is_placeholder_po = po_no.lower().startswith("pending")
    if s in ("SC Approved", "SC Pending") and sc_no:
        explicit_sc.add(sc_no)
    if s in ("PO Approved", "Finish") and po_no and not is_placeholder_po:
        explicit_po.add(po_no)


# ── Helper: get request type for SC/PO ─────────────────────────────────
def get_request_type(row, is_sc=True):
    """Determine request_type based on contract type and Internal System Number."""
    ct = row["Type of Contract"] or ""
    s = row["Status"]
    isn = row["Internal System Number （有号写号，没号写N/A）"]

    # Rule: PO under FC on GR/Finish -> ignore, keep empty
    if ct == "PO under FC" and s in ("GR Partially triggered", "Finish"):
        return ""

    # Rule: status is PO Approved or SC Approved, has Internal System Number -> call_off
    if s in ("PO Approved", "SC Approved") and is_call_off(row):
        return "call_off"

    # Otherwise map contract type
    if ct == "FC":
        return "FC"
    elif ct == "PO under FC":
        return "call_off"
    elif ct == "PO":
        return "new"
    elif ct == "Contract":
        return "new"
    return ""


# ── Generate import records ────────────────────────────────────────────
sc_records = []
po_records = []
gr_records = []

generated_sc = set(explicit_sc)
generated_po = set(explicit_po)

# Keep track of call_off SCs for reporting
call_off_scs = []


def make_sc(row):
    sc_no = row["SC No."]
    rt = get_request_type(row, is_sc=True)
    calloff_po_id = ""
    if rt == "call_off" and is_call_off(row):
        calloff_po_id = row["Internal System Number （有号写号，没号写N/A）"]
        call_off_scs.append((sc_no, calloff_po_id))

    sc_amt = row["SC amount (RMB) incl. VAT"] or sc_amount_map.get(sc_no, "")
    # Strip thousand-separator commas
    sc_amt = sc_amt.replace(",", "")

    # Map status: SC Approved/SC Pending -> approved, Finish -> finished
    s = row["Status"]
    status = "approved" if s in ("SC Approved", "SC Pending") else "finished"

    # Map requester name to user ID
    req_name = row["Requestor"]
    req_id = requester_map.get(req_name, "")
    if not req_id:
        # Try removing comma variation
        req_id = requester_map.get(req_name.replace(",", ""), "")

    # Default empty request_type to "new" for SC records
    if not rt:
        rt = "new"

    return {
        "sc_no": sc_no,
        "vendor_id": "",
        "vendor_name": row["Vendor"],
        "requester_id": req_id,
        "request_type": rt,
        "cost_center": row["Cost Center"],
        "sc_amount": sc_amt,
        "service_period_start": row["Contract from  mm/dd/yyyy"] or row["CON Date"],
        "service_period_end": row["Contract to  mm/dd/yyyy"],
        "service_scope": row["Service scope"],
        "description": row["Description"],
        "currency": "CNY",
        "asset": row["Asset (Y/N)"] or "N",
        "asset_nums": row["Asset number"],
        "internal_system_number": row["Internal System Number （有号写号，没号写N/A）"],
        "calloff_po_id": calloff_po_id,
        "status": status,
        "_source_status": row["Status"],
        "_source_row": row["_excel_row"],
    }


def make_po(row, force_po_amount=None):
    po_no = row["PO No."]
    rt = get_request_type(row, is_sc=False)

    # PO amount: try row first, then related SC Approved row, then fallback
    po_amt = force_po_amount or row["PO amount                  (RMB)"]
    if not po_amt:
        po_amt = po_amount_map.get(po_no, "")
    if not po_amt:
        # Fall back to SC amount from related row
        sc_amt = row["SC amount (RMB) incl. VAT"] or sc_amount_map.get(row["SC No."], "")
        po_amt = sc_amt
    # Strip thousand-separator commas
    po_amt = po_amt.replace(",", "")

    cf = row["Contract from  mm/dd/yyyy"] or row["CON Date"] or po_contract_from_map.get(po_no, "")
    ct = row["Contract to  mm/dd/yyyy"] or po_contract_to_map.get(po_no, "")

    # Map status: PO Approved -> active, Finish -> finished, GR Partially -> active
    s = row["Status"]
    status = "finished" if s == "Finish" else "active"

    # Map requester name to user ID
    req_name = row["Requestor"]
    req_id = requester_map.get(req_name, "")
    if not req_id:
        req_id = requester_map.get(req_name.replace(",", ""), "")

    # Default empty request_type to "new" for PO records
    if not rt:
        rt = "new"
    if not req_id:
        req_id = requester_map.get(req_name.replace(",", ""), "")

    return {
        "po_id": "",
        "sc_id": "",
        "sc_no": row["SC No."],
        "vendor_id": "",
        "vendor_name": row["Vendor"],
        "po_no": po_no,
        "request_type": rt,
        "po_amount": po_amt,
        "requester_id": req_id,
        "contract_from": cf,
        "contract_to": ct,
        "contract_no": "",
        "payment_frequency": row["Payment Frequency"],
        "contract_pos": row["Contract  Pos."],
        "contract_type": row["Type of Contract"],
        "cost_center": row["Cost Center"],
        "purchaser": row["Purchaser"],
        "status": status,
        "_source_status": row["Status"],
        "_source_row": row["_excel_row"],
    }


def make_gr(row):
    ct = row["Type of Contract"] or ""
    s = row["Status"]
    # If PO under FC on GR, ignore contract type for request_type
    rt = "" if ct == "PO under FC" else ct

    # Map status: Finish -> finished, GR Partially/Cancellation -> approved
    gr_status = "finished" if s == "Finish" else "approved"

    # Map requester name to user ID
    req_name = row["Requestor"]
    req_id = requester_map.get(req_name, "")
    if not req_id:
        req_id = requester_map.get(req_name.replace(",", ""), "")

    return {
        "po_id": "",
        "po_no": row["PO No."],
        "gr_no": row["Confirmation No."],
        "requester_id": req_id,
        "estimated_amount": row["NET GR Application Amount"] or row["GR  Application amount incl. VAT"],
        "con_value": row["CON Value"],
        "tax_rate": "",
        "goods_service_description": row["Description"],
        "confirmation_name": "",
        "is_cancellation": "Y" if s == "Cancellation" else "N",
        "delivery_from": row["Contract from  mm/dd/yyyy"] or row["CON Date"],
        "delivery_to": row["Contract to  mm/dd/yyyy"],
        "last_delivery": "",
        "status": gr_status,
        "_source_status": s,
        "_source_row": row["_excel_row"],
    }


for row in data_rows:
    if row.get("_skip"):
        continue
    s = row["Status"]
    sc_no = row["SC No."]
    po_no = row["PO No."]
    is_placeholder_po = po_no.lower().startswith("pending")

    if s in ("SC Approved", "SC Pending"):
        sc_records.append(make_sc(row))
        # Only generate PO if not placeholder AND no explicit PO row exists
        if po_no and not is_placeholder_po and po_no not in explicit_po:
            po_records.append(make_po(row))
            generated_po.add(po_no)

    elif s == "PO Approved":
        if not is_placeholder_po:
            po_records.append(make_po(row))
            if sc_no and sc_no not in explicit_sc:
                sc_records.append(make_sc(row))
                generated_sc.add(sc_no)

    elif s == "Finish":
        if not is_placeholder_po:
            sc_records.append(make_sc(row))
            po_records.append(make_po(row))
            gr_records.append(make_gr(row))
            generated_sc.add(sc_no)
            generated_po.add(po_no)

    elif s == "GR Partially triggered":
        gr_records.append(make_gr(row))
        if po_no and not is_placeholder_po and po_no not in generated_po:
            po_records.append(make_po(row))
            generated_po.add(po_no)
            if sc_no and sc_no not in generated_sc:
                sc_records.append(make_sc(row))
                generated_sc.add(sc_no)

    elif s == "Cancellation":
        gr_records.append(make_gr(row))


# ── Deduplicate ────────────────────────────────────────────────────────
def dedup(records, key_fn):
    seen = set()
    result = []
    for rec in records:
        k = key_fn(rec)
        if k not in seen:
            seen.add(k)
            result.append(rec)
        else:
            print(f"  DEDUP: {k} (from status={rec.get('_source_status','')})")
    return result


deduped_sc = dedup(sc_records, lambda r: r["sc_no"])
deduped_po = dedup(po_records, lambda r: r["po_no"])
deduped_gr = dedup(gr_records, lambda r: r["gr_no"])

# Remove _source_status / _source_row before writing
for recs in [deduped_sc, deduped_po, deduped_gr]:
    for r in recs:
        r.pop("_source_status", None)
        r.pop("_source_row", None)

print(f"\nGenerated: {len(deduped_sc)} SC, {len(deduped_po)} PO, {len(deduped_gr)} GR records")
print(f"  (before dedup: {len(sc_records)} SC, {len(po_records)} PO, {len(gr_records)} GR)")

# ── Write import CSV files ─────────────────────────────────────────────
SC_FIELDS = [
    "sc_no", "vendor_id", "vendor_name", "requester_id", "request_type",
    "cost_center", "sc_amount", "service_period_start", "service_period_end",
    "service_scope", "description", "currency", "asset", "asset_nums",
    "internal_system_number", "calloff_po_id", "status", "import_error", "error_type",
]
with open(os.path.join(BASE, "import_sc.csv"), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=SC_FIELDS, extrasaction="ignore")
    w.writeheader()
    w.writerows(deduped_sc)
print(f"Wrote {len(deduped_sc)} rows to import_sc.csv")

PO_FIELDS = [
    "po_id", "sc_id", "sc_no", "vendor_id", "vendor_name", "po_no", "request_type",
    "po_amount", "requester_id", "contract_from", "contract_to",
    "contract_no", "payment_frequency", "contract_pos", "contract_type",
    "cost_center", "purchaser", "status", "import_error", "error_type",
]
with open(os.path.join(BASE, "import_po.csv"), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=PO_FIELDS, extrasaction="ignore")
    w.writeheader()
    w.writerows(deduped_po)
print(f"Wrote {len(deduped_po)} rows to import_po.csv")

GR_FIELDS = [
    "po_id", "po_no", "gr_no", "requester_id", "estimated_amount", "con_value",
    "tax_rate", "goods_service_description", "confirmation_name",
    "delivery_from", "delivery_to", "last_delivery", "status",
    "is_cancellation", "import_error", "error_type",
]
with open(os.path.join(BASE, "import_gr.csv"), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=GR_FIELDS, extrasaction="ignore")
    w.writeheader()
    w.writerows(deduped_gr)
print(f"Wrote {len(deduped_gr)} rows to import_gr.csv")

# ══════════════════════════════════════════════════════════════════════════
# ANALYSIS
# ══════════════════════════════════════════════════════════════════════════
print("\n" + "=" * 60)
print("DATA QUALITY ANALYSIS")
print("=" * 60)

# ── Call-off SC summary ────────────────────────────────────────────────
print(f"\n1. Call-off SC records (Internal System Number present): {len(call_off_scs)}")
for sc_no, fc_po in call_off_scs:
    print(f"   SC {sc_no} -> FC PO {fc_po}")

# ── Vendor matching ────────────────────────────────────────────────────
def normalize(name):
    n = name.strip().lower()
    for ch in ".,，。、（）()":
        n = n.replace(ch, "")
    return " ".join(n.split())

xlsx_vendors = sorted(set(r["Vendor"] for r in data_rows if r["Vendor"] and not r.get("_skip")))
print(f"\n2. Unique vendors in data (after fixes): {len(xlsx_vendors)}")

wb2 = openpyxl.load_workbook(VENDOR_XLSX, data_only=True)
ws2 = wb2["Sheet1"]
db_vendors_en = {}
db_vendors_cn = {}
for r in range(2, ws2.max_row + 1):
    en = str(ws2.cell(row=r, column=2).value or "").strip()
    cn = str(ws2.cell(row=r, column=3).value or "").strip()
    if en:
        db_vendors_en[normalize(en)] = en
    if cn:
        db_vendors_cn[normalize(cn)] = cn

unmatched = []
for v in xlsx_vendors:
    nv = normalize(v)
    if nv not in db_vendors_en and nv not in db_vendors_cn:
        unmatched.append(v)

print(f"   Vendors in DB: {ws2.max_row - 1}")
print(f"   Unmatched: {len(unmatched)}")
for v in unmatched:
    print(f"     '{v}'")

# ── Request type summary ───────────────────────────────────────────────
print(f"\n3. Request types in generated SC records:")
rt_counts = defaultdict(int)
for r in deduped_sc:
    rt_counts[r["request_type"] or "(empty)"] += 1
for rt, c in sorted(rt_counts.items()):
    print(f"   '{rt}': {c}")

print(f"\n4. Request types in generated PO records:")
rt_po_counts = defaultdict(int)
for r in deduped_po:
    rt_po_counts[r["request_type"] or "(empty)"] += 1
for rt, c in sorted(rt_po_counts.items()):
    print(f"   '{rt}': {c}")

# ── PO amounts check ───────────────────────────────────────────────────
empty_po_amt = [r for r in deduped_po if not r["po_amount"]]
print(f"\n5. PO records with empty po_amount: {len(empty_po_amt)}")
for r in empty_po_amt:
    print(f"   {r['po_no']} (sc_no={r['sc_no']})")

# ── Key counts ─────────────────────────────────────────────────────────
print(f"\n6. Summary:")
print(f"   full_data.csv: {len(full_rows)} rows")
print(f"   import_sc.csv: {len(deduped_sc)} SC records")
print(f"   import_po.csv: {len(deduped_po)} PO records")
print(f"   import_gr.csv: {len(deduped_gr)} GR records")
print(f"   Cancellation GR: {len([r for r in deduped_gr if 'Cancel' in r.get('_source_status', '')])}")

print("\nDone.")
